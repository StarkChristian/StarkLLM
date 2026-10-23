"""Regressions found by a scenario test on a fresh install (invented German maintenance documents, qwen2.5vl:7b on CPU):

* a log without sentence punctuation became ONE giant chunk, so the value inside could never be retrieved
* the lexical (keyword) branch scored only the first 50 chunks of a collection
* German function words and umlaut spelling ("Zähler" vs "Zaehler") distorted the keyword ranking
* the English-only cross-encoder replaced the hybrid order and could drop the best candidate
* blank chat messages were sent to the model
"""
import os
import sys
import uuid

import pytest
from pydantic import ValidationError

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.api.chat import ChatRequest
from app.core.config import settings
from app.services import rag_service as rag_module
from app.services.rag_service import rag_service, _fold


@pytest.fixture()
def collection():
    name = f"test_{uuid.uuid4().hex[:12]}"
    col = rag_service.chroma_client.get_or_create_collection(name)
    yield col
    rag_service.chroma_client.delete_collection(name)


def _add(col, texts):
    col.add(ids=[f"c{i}" for i in range(len(texts))], documents=texts,
            metadatas=[{"source": "log.txt"} for _ in texts], embeddings=[[0.1] * 8 for _ in texts])


def test_unpunctuated_text_is_not_one_giant_chunk(monkeypatch):
    monkeypatch.setattr(rag_module.llm_service, "generate_embeddings_batch", lambda model, prompts: [[1.0, 0.0] for _ in prompts])
    text = "\n".join(f"2026-09-01T06:{i % 60:02d}:00 FT480 status ok temp {38 + i % 9} last {100 + i}" for i in range(300))
    chunks = rag_service._semantic_chunk_text(text)
    assert len(chunks) >= 20, f"{len(chunks)} chunk(s) for {len(text)} chars"
    assert max(len(c) for c in chunks) <= settings.RAG_CHUNK_SIZE * 1.1


def test_keyword_branch_scores_every_chunk(collection):
    rows = [f"2026-09-01 FT480 status ok temp {38 + i % 9} last {100 + i}" for i in range(200)]
    rows[180] = "2026-09-09 FT480 FEHLER E-57 Lichtschranke verschmutzt Zaehler 1182"
    _add(collection, rows)
    res = rag_service._keyword_query(collection, "Welcher Zähler beim Fehler E-57?", 10)
    assert res and "1182" in res[0]["content"], [r["content"][:60] for r in res[:3]]


def test_umlauts_are_folded():
    assert _fold("Zähler Straße Größe") == "zaehler strasse groesse"


def test_german_stop_words_do_not_drive_the_ranking(collection):
    _add(collection, ["Zwischen der Firma und dem Kunden wird der Vertrag geschlossen, wie der Anhang beschreibt.",
                      "Seriennummer SN-88123 des Motors wurde am Montag geprüft."])
    res = rag_service._keyword_query(collection, "Wie lautet die Seriennummer des Motors?", 5)
    assert res and "SN-88123" in res[0]["content"]


def test_rerank_never_drops_the_best_hybrid_or_vector_candidate(monkeypatch):
    cands = [{"content": f"c{i}", "metadata": {}, "id": str(i)} for i in range(10)]
    cands[4]["_vector_top"] = True

    class ReversedRanker:   # simulates a reranker that is wrong on the language of the query
        def rerank(self, request):
            return [{"id": p["id"], "score": float(i)} for i, p in enumerate(reversed(request.passages))]

    monkeypatch.setattr(rag_service, "ranker", ReversedRanker())
    top = rag_service._rerank_or_truncate(cands, "Frage", 3)
    assert len(top) == 3 and cands[0] in top and cands[4] in top


@pytest.mark.parametrize("blank", ["", "   ", "\n\t"])
def test_blank_chat_message_is_rejected(blank):
    with pytest.raises(ValidationError):
        ChatRequest(workspace_id=1, message=blank)


def test_oversized_chat_message_is_rejected():
    with pytest.raises(ValidationError):
        ChatRequest(workspace_id=1, message="x" * 100_001)


# ---- files that used to be "indexed" with nothing in the index -------------------------------------------------
@pytest.fixture()
def fake_embeddings(monkeypatch):
    monkeypatch.setattr(rag_module.llm_service, "generate_embeddings_batch", lambda model, prompts: [[1.0, 0.0] for _ in prompts])


def test_csv_rows_are_indexed_with_their_header(tmp_path):
    p = tmp_path / "teilepreise.csv"
    p.write_text("Teilenummer,Bezeichnung,Preis_EUR\nBR-2201,Umlenkrolle 80 mm,48.90\nBR-2210,Umlenkrolle 120 mm,84.50\n", encoding="utf-8")
    chunks, metas, ids = rag_service._process_file_to_chunks(str(p))
    assert chunks and "BR-2210" in " ".join(chunks) and chunks[0].startswith("Columns: Teilenummer")
    assert len(ids) == len(set(ids)) == len(metas)


def test_docx_tables_are_indexed(tmp_path, fake_embeddings):
    import docx
    d = docx.Document()
    d.add_paragraph("Reaktionszeiten nach Prioritaet:")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Prioritaet 1", "8 Stunden"
    p = tmp_path / "vertrag.docx"
    d.save(p)
    chunks, _, _ = rag_service._process_file_to_chunks(str(p))
    assert "8 Stunden" in " ".join(chunks)


def test_xlsx_rows_are_indexed(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    wb.active.append(["Teil", "Preis"])
    wb.active.append(["BR-9001", 112.0])
    p = tmp_path / "preise.xlsx"
    wb.save(p)
    chunks, metas, _ = rag_service._process_file_to_chunks(str(p))
    assert "BR-9001" in " ".join(chunks) and metas[0]["type"] == "xlsx"


def test_file_without_content_is_not_reported_as_indexed(tmp_path):
    p = tmp_path / "leer.csv"
    p.write_text("\n \n", encoding="utf-8")
    with pytest.raises(ValueError):
        rag_service._process_file_to_chunks(str(p))
